"""Training-only ingestion of explicit labels, grouped by original video bytes.

The workbook and private manifest are never consulted by ordinary inference.
Possible categories are sparse annotations, not positive labels by themselves.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
import hashlib
import io
import math
from pathlib import Path
import random
import re

from openpyxl import load_workbook

from scripts.evaluate_local import explicit_binary, point_times


_HEADERS = ("номер кейса", "таймстемп", "нарушение / без нарушения", "возможные нарушения")
_CATEGORY = re.compile(r"[a-z][a-z0-9]*(?:_[a-z0-9]+)+\Z")
_DIGEST = re.compile(r"[0-9a-f]{64}\Z")


def _blank(value) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _categories(value) -> list[str]:
    if _blank(value):
        return []
    if not isinstance(value, str):
        raise ValueError("Possible categories must contain canonical category IDs")
    values = re.split(r"[\s,;|]+", value.strip())
    if any(_CATEGORY.fullmatch(item) is None for item in values):
        raise ValueError("Possible categories must contain canonical category IDs, not descriptions")
    return sorted(set(values))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_simplified_labels(labels_path: str | Path, videos_dir: str | Path) -> dict:
    """Read an immutable source snapshot into a private, hash-grouped manifest.

    A case ID must exactly match a MOV filename stem. Explicit binary labels
    remain authoritative even when a negative row names a possible category.
    Missing categories are masked rather than converted to category negatives.
    Duplicate positive anchors are unioned as weak points, never independent
    videos or reviewed event boundaries. No workbook is edited or exported.
    """
    labels_path, videos_dir = Path(labels_path).resolve(), Path(videos_dir).resolve()
    raw = labels_path.read_bytes()
    files = sorted((path for path in videos_dir.iterdir() if path.is_file() and path.suffix.lower() == ".mov"),
                   key=lambda path: path.name)
    by_stem: dict[str, list[Path]] = defaultdict(list)
    for path in files:
        by_stem[path.stem].append(path)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    records = []
    try:
        sheet = workbook["labels"] if "labels" in workbook.sheetnames else workbook.active
        rows = iter(sheet.iter_rows())
        header = next(rows, None)
        if header is None or len(header) != len(_HEADERS):
            raise ValueError("The simplified label workbook needs its four expected columns")
        headers = [" ".join(str(cell.value or "").casefold().split()) for cell in header]
        if len(set(headers)) != len(_HEADERS) or set(headers) != set(_HEADERS):
            raise ValueError("The simplified label workbook needs its four expected columns")
        columns = [headers.index(name) for name in _HEADERS]
        for row_number, cells in enumerate(rows, 2):
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("Formula cells are not accepted as training labels")
            values = [cell.value for cell in cells]
            if all(_blank(value) for value in values):
                continue
            if len(values) != len(_HEADERS):
                raise ValueError("Every training label row needs the four expected columns")
            case_id, timestamp, verdict, possible_categories = (values[column] for column in columns)
            if not isinstance(case_id, str) or not case_id.strip():
                raise ValueError("Every training row needs an exact video stem")
            matching = by_stem.get(case_id.strip(), [])
            if len(matching) != 1:
                raise ValueError("A training case has no unique exact MOV stem match")
            try:
                label = int(explicit_binary(verdict))
            except ValueError:
                raise ValueError("Every training row needs an explicit binary label") from None
            try:
                anchors = list(point_times(timestamp))
            except ValueError:
                raise ValueError("A training row contains an invalid weak timestamp") from None
            if any(type(anchor) not in (int, float) or not math.isfinite(anchor) or anchor < 0 for anchor in anchors):
                raise ValueError("A training row contains an invalid weak timestamp")
            records.append({"video_path": str(matching[0].resolve()), "binary_label": label,
                            "observed_categories": _categories(possible_categories),
                            "anchors_sec": anchors, "source_row": row_number})
    finally:
        workbook.close()
    if not records:
        raise ValueError("The training workbook contains no labeled videos")

    hashes = {path: _sha256(Path(path)) for path in {record["video_path"] for record in records}}
    grouped: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        grouped[hashes[record["video_path"]]].append(record)
    groups = []
    for digest, members in sorted(grouped.items()):
        labels = {member["binary_label"] for member in members}
        if len(labels) != 1:
            raise ValueError("Byte-identical videos have conflicting binary labels")
        label = labels.pop()
        observed = sorted({category for member in members for category in member["observed_categories"]})
        paths = sorted({member["video_path"] for member in members})
        groups.append({"video_sha256": digest, "video_path": paths[0], "video_paths": paths,
                       "binary_label": label, "categories": observed if label else [],
                       "observed_categories": observed, "category_targets": {category: label for category in observed},
                       "category_mask": bool(observed),
                       "anchors_sec": sorted({anchor for member in members for anchor in member["anchors_sec"]}),
                       "source_rows": sorted(member["source_row"] for member in members)})
    binary_counts = Counter(group["binary_label"] for group in groups)
    return {"source_sha256": hashlib.sha256(raw).hexdigest(), "source_path": str(labels_path),
            "groups": groups, "counts": {"label_rows": len(records), "unique_video_groups": len(groups),
                                          "positive_groups": binary_counts[1], "negative_groups": binary_counts[0],
                                          "duplicate_hash_groups": sum(len(members) > 1 for members in grouped.values()),
                                          "groups_without_category_annotations": sum(not group["category_mask"] for group in groups)},
            "timestamp_semantics": "weak_point_anchors_not_event_boundaries",
            "category_semantics": "sparse_observed_categories_with_explicit_binary_targets_unlisted_categories_masked"}


def deterministic_split(groups: list[dict], seed: int = 17) -> dict[str, list[dict]]:
    """Reserve four videos per binary class for each holdout, keeping hashes whole.

    This is a new training-group holdout from an already inspected corpus, not
    evidence of untouched external generalization. Input order cannot affect it.
    Returned groups are independent copies of the supplied private records.
    """
    if type(seed) is not int:
        raise ValueError("The training split seed must be an integer")
    by_label: dict[int, list[dict]] = {0: [], 1: []}
    seen = set()
    for group in groups:
        digest, label = group.get("video_sha256"), group.get("binary_label")
        if not isinstance(digest, str) or _DIGEST.fullmatch(digest) is None:
            raise ValueError("Training groups need a valid video SHA-256")
        if digest in seen:
            raise ValueError("Training splits require one record per video hash group")
        if type(label) not in (int, bool) or label not in (0, 1):
            raise ValueError("Training groups need explicit binary labels")
        seen.add(digest)
        by_label[label].append(group)
    if any(len(members) < 9 for members in by_label.values()):
        raise ValueError("At least nine unique videos per binary class are needed for two four-video holdouts and training")
    rng = random.Random(seed)
    result = {"train": [], "validation": [], "test": []}
    for label in (0, 1):
        members = sorted(by_label[label], key=lambda group: group["video_sha256"])
        rng.shuffle(members)
        for name, chosen in (("validation", members[:4]), ("test", members[4:8]), ("train", members[8:])):
            result[name].extend(deepcopy(chosen))
    for members in result.values():
        members.sort(key=lambda group: group["video_sha256"])
    return result
