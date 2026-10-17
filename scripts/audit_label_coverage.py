"""Offline coverage audit of sparse workbook categories against the rule catalog.

This script is never imported by inference. It reports category coverage, not
detector recall, and never exports source case names, timestamps or video hashes.
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sys


def summarize_records(records: list[dict], catalog: dict, registry: dict,
                      scenarios: dict | None = None) -> dict:
    """Compare explicit category hints, preserving negative-row annotations."""
    positive, negative = Counter(), Counter()
    missing_positive = missing_negative = 0
    for record in records:
        label = record.get("binary_label")
        if type(label) is not int or label not in {0, 1}:
            raise ValueError("Coverage records require explicit binary labels")
        identifiers = set(record.get("observed_categories", []))
        if any(not isinstance(item, str) or re.fullmatch(r"[a-z][a-z0-9]*(?:_[a-z0-9]+)+", item) is None
               for item in identifiers):
            raise ValueError("Coverage categories must be canonical IDs")
        (positive if label else negative).update(identifiers)
        if not identifiers:
            if label:
                missing_positive += 1
            else:
                missing_negative += 1
    catalog_items = {item["id"]: item for item in catalog["items"]}
    registry_items = {item["id"]: item for item in registry["rules"]}
    visual_items = {item["id"]: item for item in (scenarios or {}).get("scenarios", [])}
    rows = []
    for identifier in sorted(positive.keys() | negative.keys()):
        catalog_item = catalog_items.get(identifier, {})
        registered = registry_items.get(identifier, {})
        visual = visual_items.get(identifier, {})
        rows.append({
            "type_id": identifier,
            "positive_rows": positive[identifier],
            "negative_rows": negative[identifier],
            "catalog_present": bool(catalog_item),
            "registry_present": bool(registered),
            "registry_enabled": registered.get("enabled", False) is True,
            "auto_emit": registered.get("enabled", False) is True and registered.get("auto_emit", False) is True,
            "automation": registered.get("automation", catalog_item.get("automation", "unmapped")),
            "required_checks": registered.get("required_checks", []),
            "action_checks": registered.get("action_checks", []),
            "visual_scenario_present": bool(visual),
            "visual_detection_path": bool(visual.get("visual_checks")) and visual.get("disposition") == "review_only",
            "visual_checks": visual.get("visual_checks", []),
            "visual_disposition": visual.get("disposition"),
            "visual_alias_of": visual.get("alias_of"),
        })
    legal_auto_count = sum(item.get("enabled") is True and item.get("auto_emit") is True
                           for item in registry_items.values())
    visual_count = sum(row["visual_detection_path"] for row in rows)
    return {
        "audit_version": "1.1",
        "scope": "offline_label_category_coverage_not_detection_accuracy",
        "summary": f"{visual_count}/{len(rows)} hinted types with visual detection paths; {legal_auto_count} legal automatic rules",
        "counts": {
            "label_rows": len(records),
            "positive_rows": sum(record["binary_label"] for record in records),
            "negative_rows": sum(1 - record["binary_label"] for record in records),
            "category_ids": len(rows),
            "positive_rows_without_category": missing_positive,
            "negative_rows_without_category": missing_negative,
            "hinted_types_with_visual_detection_paths": visual_count,
            "registered_visual_scenarios": len(visual_items),
            "legal_automatic_rules": legal_auto_count,
        },
        "categories": rows,
        "missing_catalog_ids": [row["type_id"] for row in rows if not row["catalog_present"]],
        "missing_enabled_registry_ids": [row["type_id"] for row in rows if not row["registry_enabled"]],
        "missing_visual_detection_path_ids": [row["type_id"] for row in rows if not row["visual_detection_path"]],
        "semantics": "Column C is the explicit binary label; column D is a sparse possible-category hint. Missing categories are unknown. Negative rows remain negative. Row counts are not independent source groups. Visual detection paths are registered review scenarios, separate from automatic legal rules. Static coverage does not establish recall or legal applicability.",
    }


def read_records(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["labels"]
        records = []
        for cells in sheet.iter_rows(min_row=2):
            if all(cell.value is None for cell in cells):
                continue
            if len(cells) != 4 or any(cell.data_type == "f" for cell in cells):
                raise ValueError("Expected four literal source-label columns")
            verdict = " ".join(str(cells[2].value or "").casefold().split())
            if verdict not in {"нарушение", "без нарушения"}:
                raise ValueError("Every row needs an explicit source binary verdict")
            categories = re.split(r"[\s,;|]+", str(cells[3].value or "").strip())
            records.append({"binary_label": int(verdict == "нарушение"),
                            "observed_categories": [item for item in categories if item]})
        return records
    finally:
        workbook.close()


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--labels", type=Path, default=root / "open/open_labels.simplified.xlsx")
    parser.add_argument("--catalog", type=Path, default=root / "rules/catalog.json")
    parser.add_argument("--registry", type=Path, default=root / "rules/registry.json")
    parser.add_argument("--scenarios", type=Path, default=root / "rules/scenarios.json")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = summarize_records(read_records(args.labels),
                               json.loads(args.catalog.read_text(encoding="utf-8")),
                               json.loads(args.registry.read_text(encoding="utf-8")),
                               json.loads(args.scenarios.read_text(encoding="utf-8")))
    report["source_workbook_sha256"] = hashlib.sha256(args.labels.read_bytes()).hexdigest()
    serialized = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        if args.output.resolve() in {args.labels.resolve(), args.catalog.resolve(), args.registry.resolve(), args.scenarios.resolve()}:
            parser.error("The report cannot overwrite a source file")
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(serialized, encoding="utf-8")
    else:
        sys.stdout.write(serialized)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
