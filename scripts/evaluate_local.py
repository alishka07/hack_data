"""Aggregate exploratory evaluation of local completed jobs; never print video content.

This script is evaluation-only. Workbook names/point labels never reach inference.
Original and uploaded bytes are joined by SHA-256, and byte-identical originals
remain one group. A running job never substitutes for a completed result.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, time as datetime_time, timedelta, timezone
from functools import lru_cache
import hashlib
import io
import json
import math
import os
from pathlib import Path
import random
import re
import shutil
import sqlite3
import subprocess

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PREVIOUS_AUDIT_SHA256 = "a6e80004dde165c9294acc5cae165ab5df9aa11d5dc9e0d6b50e313da7701e4e"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def point_times(value) -> tuple[float, ...]:
    """Extract weak point times locally; surrounding annotation text is never exported."""
    if value is None or isinstance(value, str) and not value.strip():
        return ()
    if isinstance(value, datetime_time):
        return (value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6,)
    if isinstance(value, timedelta):
        seconds = value.total_seconds()
        if seconds < 0:
            raise ValueError("Negative timestamp")
        return (seconds,)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        seconds = float(value)
        if not math.isfinite(seconds) or seconds < 0:
            raise ValueError("Invalid timestamp")
        return (seconds,)
    if not isinstance(value, str):
        raise ValueError("Unsupported timestamp format")
    def clock(token):
        parts = [float(part) for part in token.split(":")]
        if len(parts) > 1 and any(part >= 60 for part in parts[1:]):
            raise ValueError("Invalid clock timestamp")
        return sum(part * 60 ** i for i, part in enumerate(reversed(parts)))

    tokens = [token.strip() for token in re.split(r"[;,\n]+", value.strip()) if token.strip()]
    if all(re.fullmatch(r"\d+(?:\.\d+)?(?::\d+(?:\.\d+)?){0,2}", token) for token in tokens):
        return tuple(sorted(set(clock(token) for token in tokens)))
    if re.search(r"(?<!\w)-\s*\d", value):
        raise ValueError("Ambiguous negative timestamp")
    result = [clock(match.group()) for match in re.finditer(r"(?<![\d:])\d+(?::\d{1,2}){1,2}(?![\d:])", value)]
    units = list(re.finditer(
        r"(?<![\w\d])(?P<number>\d+(?:[.,]\d+)?)\s*(?P<unit>"
        r"мин(?:ут(?:а|ы|у|е|ах)?)?|minutes?|min|сек(?:унд(?:а|ы|у|е|ах)?)?|seconds?|sec|s)\.?(?!\w)",
        value, re.I))
    index = 0
    while index < len(units):
        match = units[index]
        number = float(match["number"].replace(",", "."))
        minute = match["unit"].casefold().startswith(("мин", "min"))
        if minute and index + 1 < len(units) and not units[index + 1]["unit"].casefold().startswith(("мин", "min")):
            second = units[index + 1]
            if not re.fullmatch(r"[\s,.;:]*(?:и[\s,.;:]*)?", value[match.end():second.start()]):
                raise ValueError("Ambiguous minutes/seconds relationship")
            seconds = float(second["number"].replace(",", "."))
            if seconds >= 60:
                raise ValueError("Invalid minutes/seconds timestamp")
            result.append(number * 60 + seconds)
            index += 2
        else:
            result.append(number * 60 if minute else number)
            index += 1
    if not result:
        raise ValueError("No timestamp")
    return tuple(sorted(set(result)))


def is_blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def explicit_binary(value):
    """Use a strict binary vocabulary; never classify arbitrary descriptive text."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = " ".join(value.casefold().split())
        if normalized in {"нарушение", "violation", "yes", "true", "1", "да"}:
            return True
        if normalized in {"без нарушения", "без нарушений", "нет нарушения", "нарушений нет", "no_violation", "no violation", "no", "false", "0", "нет"}:
            return False
    raise ValueError("Unrecognized explicit binary label")


def read_reference(labels_path: Path, videos_dir: Path, file_hash=sha256, snapshot_dir: Path | None = None):
    """Take one immutable workbook snapshot and return derived metadata only.

    The third return value is private row metadata for local audit files. It has
    names, hashes and extracted point times, but no annotation/description text.
    """
    raw = labels_path.read_bytes()
    if snapshot_dir is not None:
        snapshot_dir.mkdir(parents=True, exist_ok=True)
        snapshot = snapshot_dir / (hashlib.sha256(raw).hexdigest() + ".xlsx")
        if not snapshot.exists():
            snapshot.write_bytes(raw)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    files = [p for p in videos_dir.iterdir() if p.is_file() and p.suffix.lower() == ".mov"]
    originals = {path.name: path for path in files}
    stems = defaultdict(list)
    for path in files:
        stems[path.stem].append(path)
    original_hashes = {path: file_hash(path) for path in files}
    groups, private_rows = defaultdict(list), []
    counts = Counter()
    try:
        sheet = workbook["labels"] if "labels" in workbook.sheetnames else workbook.active
        formula_sheet = formulas[sheet.title]
        rows = list(sheet.iter_rows(values_only=True))
        headers = list(rows[0]) if rows else []
        normalized = [" ".join(str(value or "").casefold().split()) for value in headers]
        filename_column = next((i for i, value in enumerate(normalized) if value in {"наименование файла", "file", "filename", "file_name"}), 0)
        timestamp_column = next((i for i, value in enumerate(normalized) if value in {"таймстемп", "timestamp", "timestamps", "time", "timecode", "время"}), 1)
        verdict_columns = [i for i, value in enumerate(normalized) if value in {"verdict", "label", "binary label", "violation", "нарушение", "нарушение / без нарушения"}]
        if len(verdict_columns) > 1:
            raise ValueError("Multiple explicit verdict columns require manual schema review")
        verdict_column = verdict_columns[0] if verdict_columns else None
        input_columns = {filename_column, timestamp_column} | ({verdict_column} if verdict_column is not None else set())
        formula_cells = {(cell.row, cell.column - 1) for row in formula_sheet.iter_rows() for cell in row if cell.data_type == "f"}
        typed = [{"column": get_column_letter(i + 1), "header": headers[i],
                  "types": dict(Counter(type(row[i] if i < len(row) else None).__name__ for row in rows[1:])),
                  "nonblank_rows": sum(not is_blank(row[i] if i < len(row) else None) for row in rows[1:])}
                 for i in range(len(headers))]
        vocab = set()
        for row_number, row in enumerate(rows[1:], 2):
            get = lambda column: row[column] if column is not None and column < len(row) else None
            name_value = get(filename_column)
            if is_blank(name_value):
                if any(not is_blank(value) for value in row):
                    counts["rows_without_filename"] += 1
                continue
            counts["rows"] += 1
            name = str(name_value).strip()
            original = originals.get(name)
            if original is None and len(stems.get(name, [])) == 1:
                original = stems[name][0]
            timestamp_value, verdict_value = get(timestamp_column), get(verdict_column)
            has_annotation = not is_blank(timestamp_value)
            counts["timestamp_nonblank_rows" if has_annotation else "timestamp_blank_rows"] += 1
            issues = []
            try:
                points = point_times(timestamp_value)
                if has_annotation:
                    counts["parsed_timestamp_rows"] += 1
                    counts["extracted_point_count"] += len(points)
            except ValueError:
                points = None
                counts["unparsed_timestamp_rows"] += 1
                issues.append("timestamp_unparsed")
            if not is_blank(verdict_value):
                try:
                    binary = explicit_binary(verdict_value)
                    counts["explicit_positive_rows" if binary else "explicit_negative_rows"] += 1
                    vocab.add(str(verdict_value).strip().casefold())
                    if not binary and points:
                        binary = None
                        counts["row_binary_timestamp_conflicts"] += 1
                        issues.append("negative_verdict_with_positive_timestamp")
                except ValueError:
                    binary = None
                    counts["invalid_explicit_verdict_rows"] += 1
                    issues.append("unrecognized_binary_verdict")
                source = "explicit_verdict"
            else:
                binary = bool(points) if points is not None else None
                source = "timestamp_fallback"
                counts["fallback_rows"] += 1
            if any((row_number, column) in formula_cells and is_blank(get(column)) for column in input_columns):
                binary = None
                issues.append("input_formula_has_no_cached_value")
                counts["uncached_formula_rows"] += 1
            if binary is True and not points:
                points = None
                counts["positive_rows_without_point_times"] += 1
            counts["positive_rows" if binary is True else "negative_rows" if binary is False else "invalid_binary_rows"] += 1
            if original is None:
                counts["unmatched_filename_rows"] += 1
                digest = None
            else:
                digest = original_hashes[original]
            record = {"row": row_number, "filename": name, "sha256": digest, "binary": binary,
                      "points": points, "binary_source": source, "issues": issues,
                      "video_path": str(original.resolve()) if original is not None else None}
            private_rows.append(record)
            if digest is not None:
                groups[digest].append(record)
        group_counts = Counter({"duplicate_hash_groups": 0, "disagreeing_timestamp_groups": 0,
                                "conflicting_binary_groups": 0, "positive_groups": 0,
                                "negative_groups": 0, "invalid_binary_groups": 0})
        for records in groups.values():
            labels = {record["binary"] for record in records if record["binary"] is not None}
            point_sets = {record["points"] for record in records if record["points"] is not None}
            group_counts["duplicate_hash_groups"] += len(records) > 1
            group_counts["disagreeing_timestamp_groups"] += len(point_sets) > 1
            group_counts["conflicting_binary_groups"] += len(labels) > 1
            valid = len(labels) == 1 and all(record["binary"] is not None for record in records)
            if valid:
                group_counts["positive_groups" if next(iter(labels)) else "negative_groups"] += 1
            else:
                group_counts["invalid_binary_groups"] += 1
        summary = {"workbook_sha256": hashlib.sha256(raw).hexdigest(), "sheet": sheet.title,
                   "range": f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}", "columns": typed,
                   "filename_column": get_column_letter(filename_column + 1),
                   "timestamp_column": get_column_letter(timestamp_column + 1),
                   "explicit_verdict_column": get_column_letter(verdict_column + 1) if verdict_column is not None else None,
                   "explicit_binary_vocabulary": sorted(vocab),
                   "precedence": "Explicit nonblank binary verdict is authoritative; conflicting positive timestamps invalidate a negative verdict. Blank=negative fallback applies only when the explicit verdict is absent.",
                   "label_rows": counts["rows"], "positive_rows": counts["positive_rows"], "negative_rows": counts["negative_rows"],
                   "original_video_files": len(files), "original_unique_hash_groups": len(set(original_hashes.values())),
                   "matched_unique_hash_groups": len(groups), **dict(group_counts),
                   "row_counts": dict(counts), "formula_input_cell_count": sum(column in input_columns for _, column in formula_cells),
                   "ground_truth_human_validated": False,
                   "timestamp_semantics": "Weak point anchors extracted locally; not reviewed event boundaries or category labels",
                   "privacy": "No scene descriptions, original annotation text, frames, filenames or row-level labels in this aggregate audit"}
        return groups, summary, private_rows
    finally:
        workbook.close()
        formulas.close()


def match_points(reference: tuple[float, ...], predicted: tuple[float, ...], tolerance: float) -> tuple[float, ...]:
    """One-to-one ordered matching: maximize match count, then minimize timing error."""
    reference, predicted = tuple(sorted(reference)), tuple(sorted(predicted))

    @lru_cache(None)
    def best(i, j):
        if i == len(reference) or j == len(predicted):
            return ()
        choices = [best(i + 1, j), best(i, j + 1)]
        distance = abs(reference[i] - predicted[j])
        if distance <= tolerance:
            choices.append((distance,) + best(i + 1, j + 1))
        return max(choices, key=lambda distances: (len(distances), -sum(distances)))

    return best(0, 0)


def ratio(numerator, denominator):
    return round(numerator / denominator, 6) if denominator else None


def candidate_metadata(raw):
    """Whitelist numeric timing/version fields; never return arbitrary detail strings."""
    version, proposals = "unknown", None
    try:
        details = json.loads(raw) if raw else {}
        candidate_version = details.get("versions", {}).get("pipeline")
        if isinstance(candidate_version, str) and re.fullmatch(r"\d+(?:\.\d+)*", candidate_version):
            version = candidate_version
        events = details.get("candidate_events")
        if isinstance(events, list):
            values = [event["timecode_sec"] for event in events]
            if all(isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value >= 0 for value in values):
                proposals = tuple(float(value) for value in values)
    except (TypeError, ValueError, KeyError, AttributeError):
        pass
    return {"pipeline_version": version, "candidate_points": proposals}


def prepare_validation(labels_path: Path, videos_dir: Path, audit_dir: Path, seed: int = 17):
    """Create private exploratory partitions; this corpus was already inspected."""
    groups, reference, _ = read_reference(labels_path, videos_dir, snapshot_dir=audit_dir)
    classified = {}
    eligible = {True: [], False: []}
    for digest, records in sorted(groups.items()):
        labels = {record["binary"] for record in records}
        points = {record["points"] for record in records if record["points"] is not None}
        valid = len(labels) == 1 and None not in labels
        label = next(iter(labels)) if valid else None
        disagreement = len(points) > 1
        classified[digest] = {"sha256": digest, "binary": label,
                              "video_paths": sorted({record["video_path"] for record in records}),
                              "filenames": sorted({record["filename"] for record in records}),
                              "timestamp_disagreement": disagreement,
                              "point_times_sec": sorted(next(iter(points))) if len(points) == 1 else None}
        if valid and not disagreement and all(record["points"] is not None for record in records):
            eligible[label].append(digest)
    generator = random.Random(seed)
    validation = []
    for label in (True, False):
        generator.shuffle(eligible[label])
        if len(eligible[label]) < 4:
            raise ValueError("Insufficient unambiguous groups for balanced 4-positive/4-negative validation")
        validation.extend(eligible[label][:4])
    validation = set(validation)
    for digest, group in classified.items():
        group["split"] = "validation" if digest in validation else "development"
    ffprobe = os.environ.get("FORTE_FFPROBE", "ffprobe")
    longest = None
    for digest, group in classified.items():
        try:
            process = subprocess.run([ffprobe, "-v", "error", "-show_entries", "format=duration", "-of", "json", group["video_paths"][0]],
                                     capture_output=True, check=True, timeout=15)
            duration = float(json.loads(process.stdout)["format"]["duration"])
            if not math.isfinite(duration) or duration <= 0:
                raise ValueError("Invalid duration")
        except (OSError, subprocess.SubprocessError, ValueError, KeyError, TypeError):
            raise ValueError("A local video duration could not be probed for validation selection") from None
        if longest is None or duration > longest["duration_sec"]:
            longest = {"sha256": digest, "video_path": group["video_paths"][0], "duration_sec": duration,
                       "split": group["split"], "binary": group["binary"]}
    samples = []
    for label in (True, False):
        chosen = [digest for digest in sorted(validation) if classified[digest]["binary"] is label][:2]
        for digest in chosen:
            samples.append({"sha256": digest, "video_path": classified[digest]["video_paths"][0],
                            "binary": label, "split": "validation"})
    caveat = "Exploratory development validation: the corpus was previously inspected; not an untouched holdout or evidence of generalization"
    summary = {"seed": seed, "workbook_sha256": reference["workbook_sha256"], "total_groups": len(classified),
               "development_groups": sum(group["split"] == "development" for group in classified.values()),
               "validation_groups": len(validation), "validation_positive_groups": 4, "validation_negative_groups": 4,
               "timestamp_disagreement_groups_development_only": sum(group["timestamp_disagreement"] and group["split"] == "development" for group in classified.values()),
               "selected_validation_samples": len(samples), "caveat": caveat}
    audit_dir.mkdir(parents=True, exist_ok=True)
    (audit_dir / "groups.json").write_text(json.dumps({"summary": summary, "groups": list(classified.values())}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (audit_dir / "validation-selection.json").write_text(json.dumps({"summary": summary, "validation_samples": samples, "longest_sample": longest}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


def evaluate(labels_path: Path, videos_dir: Path, db_path: Path, tolerance: float = 3.0,
             audit_dir: Path | None = None, aggregate_audit_path: Path | None = None,
             pipeline_version: str | None = None) -> dict:
    if not math.isfinite(tolerance) or tolerance < 0:
        raise ValueError("Timestamp tolerance must be finite and nonnegative")
    if pipeline_version is not None and not re.fullmatch(r"\d+(?:\.\d+)*", pipeline_version):
        raise ValueError("Pipeline version filter must be an exact numeric version")
    hash_cache = {}

    def file_hash(path):
        resolved = path.resolve()
        stat = resolved.stat()
        key = (resolved, stat.st_size, stat.st_mtime_ns)
        if key not in hash_cache:
            hash_cache[key] = sha256(resolved)
        return hash_cache[key]

    counts = Counter()
    labels_by_hash, reference, private_rows = read_reference(labels_path, videos_dir, file_hash, audit_dir)
    if audit_dir is not None:
        (audit_dir / "restored-audit.json").write_text(json.dumps({"summary": reference, "rows": private_rows}, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    if aggregate_audit_path is not None:
        aggregate_audit_path.parent.mkdir(parents=True, exist_ok=True)
        aggregate_audit_path.write_text(json.dumps(reference, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")

    jobs_by_hash = defaultdict(list)
    with sqlite3.connect(db_path.resolve().as_uri() + "?mode=ro", uri=True, timeout=10) as connection:
        connection.row_factory = sqlite3.Row
        jobs = connection.execute("SELECT id,status,video_path,created_at FROM jobs ORDER BY created_at DESC,id DESC").fetchall()
        for job in jobs:
            path = Path(job["video_path"])
            try:
                digest = file_hash(path)
            except OSError:
                counts["jobs_with_missing_upload"] += 1
                continue
            if digest not in labels_by_hash:
                counts["jobs_not_matching_reference_bytes"] += 1
                continue
            jobs_by_hash[digest].append(job)

        completed = {}
        completed_metadata = {}
        invalid_completed = set()
        has_details = "details_json" in {row[1] for row in connection.execute("PRAGMA table_info(jobs)")}
        for digest, grouped_jobs in jobs_by_hash.items():
            latest, stored, metadata = None, None, None
            for candidate in grouped_jobs:
                if candidate["status"] != "done":
                    continue
                stored_candidate = connection.execute("SELECT result_json" + (",details_json" if has_details else "") + " FROM jobs WHERE id=?", (candidate["id"],)).fetchone()
                candidate_info = candidate_metadata(stored_candidate[1] if has_details else None)
                if pipeline_version is not None and candidate_info["pipeline_version"] != pipeline_version:
                    counts["done_jobs_excluded_by_pipeline_filter"] += 1
                    continue
                latest, stored, metadata = candidate, stored_candidate, candidate_info
                break
            if latest is None:
                continue
            raw = stored[0]
            try:
                result = json.loads(raw)
                verdict = result["verdict"]
                events = result["violations"]
                if verdict not in {"violation", "no_violation"} or not isinstance(events, list) or bool(events) != (verdict == "violation"):
                    raise ValueError("Invalid completed result")
                timestamps = []
                for event in events:
                    value = event["timecode_sec"]
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                        raise ValueError("Invalid completed event timestamp")
                    timestamps.append(float(value))
                completed[digest] = (verdict == "violation", tuple(timestamps))
            except (TypeError, ValueError, KeyError):
                invalid_completed.add(digest)
                continue
            completed_metadata[digest] = metadata

    confusion = Counter({"true_positive": 0, "false_positive": 0, "true_negative": 0, "false_negative": 0})
    status_counts = Counter({"completed": 0, "invalid_completed": 0, "processing": 0, "failed": 0, "other_pipeline_only": 0, "not_submitted": 0})
    localization_counts = Counter({"eligible_completed_groups": 0, "positive_reference_groups": 0,
                                   "reference_points": 0, "predicted_points": 0,
                                   "completed_groups_excluded_for_disagreement": 0,
                                   "completed_groups_excluded_for_missing_timestamps": 0})
    distances = []
    proposal_counts = Counter({"eligible_completed_groups": 0, "groups_without_candidate_metadata": 0,
                               "positive_reference_groups": 0, "reference_points": 0, "proposal_points": 0})
    proposal_distances = []
    binary_excluded_completed = 0
    for digest, records in labels_by_hash.items():
        binary_labels = {record["binary"] for record in records if record["binary"] is not None}
        point_sets = {record["points"] for record in records if record["points"] is not None}
        timestamp_conflict = len(point_sets) > 1
        grouped_jobs = jobs_by_hash[digest]
        if digest in completed:
            status_counts["completed"] += 1
        elif digest in invalid_completed:
            status_counts["invalid_completed"] += 1
        elif any(job["status"] == "processing" for job in grouped_jobs):
            status_counts["processing"] += 1
        elif any(job["status"] == "failed" for job in grouped_jobs):
            status_counts["failed"] += 1
        elif any(job["status"] == "done" for job in grouped_jobs):
            status_counts["other_pipeline_only"] += 1
        else:
            status_counts["not_submitted"] += 1
        if digest not in completed:
            continue
        if len(binary_labels) != 1 or any(record["binary"] is None for record in records):
            binary_excluded_completed += 1
            continue
        truth = next(iter(binary_labels))
        prediction, predicted_points = completed[digest]
        confusion[("true_positive" if prediction else "false_negative") if truth else ("false_positive" if prediction else "true_negative")] += 1
        if timestamp_conflict:
            localization_counts["completed_groups_excluded_for_disagreement"] += 1
            continue
        if any(record["points"] is None for record in records):
            localization_counts["completed_groups_excluded_for_missing_timestamps"] += 1
            continue
        reference_points = records[0]["points"]
        localization_counts["eligible_completed_groups"] += 1
        localization_counts["positive_reference_groups"] += bool(reference_points)
        localization_counts["reference_points"] += len(reference_points)
        localization_counts["predicted_points"] += len(predicted_points)
        distances.extend(match_points(reference_points, predicted_points, tolerance))
        proposals = completed_metadata[digest]["candidate_points"]
        if proposals is None:
            proposal_counts["groups_without_candidate_metadata"] += 1
        else:
            proposal_counts["eligible_completed_groups"] += 1
            proposal_counts["positive_reference_groups"] += bool(reference_points)
            proposal_counts["reference_points"] += len(reference_points)
            proposal_counts["proposal_points"] += len(proposals)
            proposal_distances.extend(match_points(reference_points, proposals, tolerance))

    tp, fp, tn, fn = (confusion[key] for key in ("true_positive", "false_positive", "true_negative", "false_negative"))
    current_sha = reference["workbook_sha256"]
    warnings = [
        "Exploratory development-only results; clips were not held out and do not establish generalization.",
        "Ground truth has not been human-validated. Point annotations do not define event intervals or event categories.",
        "The service covers four reviewed visual rule categories; broad workbook violation labels may describe unsupported events.",
        "Category accuracy, evidence relevance and legal applicability require human review and are not scored here.",
        "Only the latest completed job per byte-identical video group is evaluated; active jobs are ignored as predictions.",
    ]
    if current_sha != PREVIOUS_AUDIT_SHA256:
        warnings.insert(0, "Workbook bytes differ from the original planning audit. The immutable current snapshot is evaluated; prior all-blank-reference results are superseded.")
    if reference["positive_rows"] == 0:
        warnings.insert(0, "The current reference has no valid positive labels, so positive recall and localization cannot be measured.")
    if reference["invalid_binary_groups"]:
        warnings.insert(0, "Conflicting or unrecognized binary labels are excluded, never silently overridden.")
    positive_reference = localization_counts["reference_points"]
    matched = len(distances)
    total_groups = len(labels_by_hash)
    return {
        "purpose": "Exploratory local development evaluation; not held-out accuracy or generalization evidence",
        "pipeline_version_filter": pipeline_version,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "warnings": warnings,
        "reference": {**reference, "previous_audit_sha256": PREVIOUS_AUDIT_SHA256,
                      "changed_since_prior_audit": current_sha != PREVIOUS_AUDIT_SHA256,
                      "blank_timestamp_fallback": "no_violation only when explicit binary verdict is absent"},
        "coverage": {"completed_groups": status_counts["completed"],
                     "unevaluated_groups": total_groups - status_counts["completed"],
                     "group_status_disjoint": dict(status_counts),
                     "database_jobs_seen": len(jobs), "jobs_with_missing_upload": counts["jobs_with_missing_upload"],
                     "jobs_not_matching_reference_bytes": counts["jobs_not_matching_reference_bytes"],
                     "done_jobs_excluded_by_pipeline_filter": counts["done_jobs_excluded_by_pipeline_filter"],
                     "completed_pipeline_versions": dict(Counter(metadata["pipeline_version"] for metadata in completed_metadata.values()))},
        "exploratory_binary": {"evaluated_groups": tp + fp + tn + fn,
                               "completed_groups_excluded_for_label_issues": binary_excluded_completed,
                               "reference_positive_groups": reference["positive_groups"],
                               "reference_negative_groups": reference["negative_groups"],
                               "confusion": dict(confusion), "precision": ratio(tp, tp + fp),
                               "recall": ratio(tp, tp + fn), "specificity": ratio(tn, tn + fp),
                               "f1": ratio(2 * tp, 2 * tp + fp + fn)},
        "exploratory_localization": {"status": "available" if positive_reference else "unavailable_no_evaluated_positive_reference_points",
                                     "tolerance_sec": tolerance, **dict(localization_counts),
                                     "matched_points": matched,
                                     "unmatched_reference_points": positive_reference - matched,
                                     "unmatched_predicted_points": localization_counts["predicted_points"] - matched,
                                     "precision": ratio(matched, localization_counts["predicted_points"]) if positive_reference else None,
                                     "recall": ratio(matched, positive_reference),
                                     "mean_absolute_error_sec": ratio(sum(distances), matched)},
        "exploratory_candidate_localization": {"status": "available" if proposal_counts["positive_reference_groups"] else "unavailable_no_evaluated_positive_candidate_metadata",
                                                "purpose": "Proposal timing before the legal evidence gate; not a violation or legality score",
                                                "tolerance_sec": tolerance, **dict(proposal_counts),
                                                "matched_points": len(proposal_distances),
                                                "unmatched_reference_points": proposal_counts["reference_points"] - len(proposal_distances),
                                                "recall": ratio(len(proposal_distances), proposal_counts["reference_points"]),
                                                "mean_absolute_error_sec": ratio(sum(proposal_distances), len(proposal_distances))},
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--labels", type=Path, default=ROOT / "open" / "open_labels.xlsx")
    parser.add_argument("--videos-dir", type=Path, default=ROOT / "open" / "videos")
    data_dir = Path(os.environ.get("FORTE_DATA_DIR", ROOT / "data"))
    parser.add_argument("--db", type=Path, default=Path(os.environ.get("FORTE_DB_PATH", data_dir / "jobs.sqlite")))
    parser.add_argument("--output", type=Path, default=ROOT / "docs" / "evaluation.json")
    parser.add_argument("--audit-dir", type=Path, default=ROOT / "data" / "labels")
    parser.add_argument("--aggregate-audit-output", type=Path, default=ROOT / "docs" / "dataset-audit.json")
    parser.add_argument("--tolerance-sec", type=float, default=3.0)
    parser.add_argument("--pipeline-version", help="Use only the latest completed job with this exact pipeline version per video hash")
    parser.add_argument("--prepare-validation", action="store_true", help="Write only private deterministic exploratory validation manifests; preserve evaluation reports")
    args = parser.parse_args()
    if args.prepare_validation:
        print(json.dumps(prepare_validation(args.labels, args.videos_dir, args.audit_dir), ensure_ascii=True))
        return
    report = evaluate(args.labels, args.videos_dir, args.db, args.tolerance_sec, args.audit_dir, args.aggregate_audit_output, args.pipeline_version)
    if args.pipeline_version is not None and not report["coverage"]["completed_groups"]:
        report["warnings"].insert(0, "No completed jobs match the exact pipeline filter. Existing evaluation report was preserved.")
        print(json.dumps(report, ensure_ascii=True, allow_nan=False))
        return
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.output.exists():
        previous = json.loads(args.output.read_text(encoding="utf-8"))
        if previous.get("reference", {}).get("workbook_sha256") == "6a5c0feddc76ed2467dcc6bfc8673402d69e4e1270c1ef1ad6643f6971b99131":
            archive = args.output.with_name("evaluation-allblank-superseded.json")
            shutil.copyfile(args.output, archive)
            report["supersedes"] = {"report": archive.name, "reason": "Restored explicit binary labels replace the incomplete all-blank reference"}
    archive = args.output.with_name("evaluation-allblank-superseded.json")
    if archive.exists():
        report["supersedes"] = {"report": archive.name, "reason": "Restored explicit binary labels replace the incomplete all-blank reference"}
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True, allow_nan=False))


if __name__ == "__main__":
    main()
