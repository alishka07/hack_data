"""Training-only provenance audit. Never import this module from inference.

Filename families prevent accidental augmentation leakage, but do not identify
sources or labels. Perceptual matches are review suggestions, never supervision.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import io
import json
import math
from pathlib import Path
import re

import cv2
import numpy as np
from openpyxl import load_workbook

from scripts.evaluate_local import explicit_binary, point_times
from worker.training_data import deterministic_split


SCHEMA_VERSION = "forte-augmentation-audit/1"
FRACTIONS = (0.1, 0.3, 0.5, 0.7, 0.9)
VIDEO_SUFFIXES = {".mov", ".mp4", ".mkv", ".avi", ".m4v", ".webm"}
_HASH = re.compile(r"[0-9a-f]{64}\Z")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def augmentation_family(path: Path) -> tuple[str, str | None]:
    """A naming hint can group derivatives, but never maps cashier_NN to a case."""
    stem = path.stem
    if "__" not in stem:
        return stem, None
    family, transform = stem.split("__", 1)
    return family, transform or None


def read_explicit_workbook(path: Path, videos_dir: Path) -> dict:
    """Read columns A/B/C strictly; blank/unknown C is unknown, never negative.

    D and later columns are deliberately excluded from supervision. Weak point
    anchors do not establish categories, visibility, or event boundaries.
    """
    raw = path.read_bytes()
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    files = [p for p in videos_dir.iterdir() if p.is_file() and p.suffix.lower() in VIDEO_SUFFIXES]
    names = defaultdict(list)
    for video in files:
        names[video.name].append(video)
        names[video.stem].append(video)
    records, issues = [], Counter()
    try:
        sheet = workbook["labels"] if "labels" in workbook.sheetnames else workbook.active
        rows = iter(sheet.iter_rows())
        header = next(rows, ())
        if len(header) < 3 or " ".join(str(header[2].value or "").casefold().split()) != "нарушение / без нарушения":
            raise ValueError("Column C must be the explicit binary verdict")
        for row_number, row in enumerate(rows, 2):
            values = [cell.value for cell in row]
            if not any(value is not None and str(value).strip() for value in values):
                continue
            errors = []
            if len(row) < 3 or any(cell.data_type == "f" for cell in row[:3]):
                errors.append("missing_or_formula_input")
            key = str(values[0] or "").strip()
            matches = names.get(key, [])
            if len(matches) != 1:
                errors.append("no_unique_exact_filename_or_stem")
            try:
                label = int(explicit_binary(values[2])) if len(values) > 2 else None
            except ValueError:
                label = None
                errors.append("missing_or_invalid_explicit_column_c")
            try:
                anchors = list(point_times(values[1])) if len(values) > 1 else []
            except ValueError:
                anchors = []
                errors.append("invalid_weak_anchor")
            if errors:
                label = None
            issues.update(errors)
            records.append({"row": row_number, "video_path": str(matches[0].resolve()) if len(matches) == 1 else None,
                            "binary_label": label, "weak_anchors_sec": anchors, "issues": errors})
    finally:
        workbook.close()
    return {"path": str(path.resolve()), "sha256": hashlib.sha256(raw).hexdigest(),
            "column_c_authoritative": True, "rows": records, "issues": dict(issues)}


def inspect_video(path: Path) -> tuple[dict, dict | None]:
    """Hash every byte and decode five temporal samples, with both orientations.

    OpenCV duration/frame estimates are audit metadata, not authoritative source
    PTS and are never used to generate training event boundaries.
    """
    before = path.stat()
    record = {"path": str(path.resolve()), "sha256": sha256_file(path), "size_bytes": before.st_size,
              "decode_status": "failed", "metadata_semantics": "opencv_estimates_not_source_pts"}
    cap = cv2.VideoCapture(str(path))
    signature = None
    try:
        fps, count = cap.get(cv2.CAP_PROP_FPS), cap.get(cv2.CAP_PROP_FRAME_COUNT)
        if not cap.isOpened() or not math.isfinite(fps) or fps <= 0 or not math.isfinite(count) or count < 1:
            return record, None
        record.update({"fps_estimate": fps, "frame_count_estimate": int(count), "duration_sec_estimate": count / fps,
                       "width": int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)), "height": int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))})
        vectors, flipped, hashes, flipped_hashes, photometric = [], [], [], [], []
        for fraction in FRACTIONS:
            cap.set(cv2.CAP_PROP_POS_FRAMES, round(fraction * (count - 1)))
            ok, frame = cap.read()
            if not ok:
                return record, None
            luma = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
            photometric.append({"mean_luma_0_255": round(float(luma.mean()), 4),
                                "std_luma": round(float(luma.std()), 4),
                                "near_clipped_fraction": round(float(((luma <= 2) | (luma >= 253)).mean()), 6)})
            for mirrored, target_vectors, target_hashes in ((False, vectors, hashes), (True, flipped, flipped_hashes)):
                image = cv2.flip(frame, 1) if mirrored else frame
                gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
                small = cv2.resize(gray, (64, 36), interpolation=cv2.INTER_AREA).astype(np.float32)
                small = cv2.GaussianBlur(small, (3, 3), 0)
                vector = small.reshape(-1) - small.mean()
                norm = float(np.linalg.norm(vector))
                vector = vector / norm if norm > 1e-6 else np.zeros_like(vector)
                target_vectors.append(vector)
                dct = cv2.dct(cv2.resize(gray, (32, 32), interpolation=cv2.INTER_AREA).astype(np.float32))[:8, :8].reshape(-1)
                bits = dct > np.median(dct[1:])
                bits[0] = False
                target_hashes.append(int.from_bytes(np.packbits(bits).tobytes(), "big"))
        record["decode_status"] = "sampled"
        record["sample_fractions"] = list(FRACTIONS)
        record["perceptual_hashes"] = [f"{value:016x}" for value in hashes]
        record["photometric_samples"] = photometric
        signature = {"normal": np.stack(vectors), "flipped": np.stack(flipped),
                     "hashes": hashes, "flipped_hashes": flipped_hashes}
    finally:
        cap.release()
    after = path.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise ValueError("A source changed during the audit")
    return record, signature


def rank_source_candidates(clip: dict, signature: dict | None, sources: list[dict], signatures: dict) -> list[dict]:
    """Return conservative sequence similarity hints; never return a label."""
    if signature is None:
        return []
    ranked = []
    for source in sources:
        source_signature = signatures.get(source["sha256"])
        if source_signature is None:
            continue
        duration_error = abs(clip["duration_sec_estimate"] - source["duration_sec_estimate"])
        if duration_error > max(0.25, 0.01 * source["duration_sec_estimate"]):
            continue
        orientations = []
        for orientation, hash_key in (("normal", "hashes"), ("flipped", "flipped_hashes")):
            correlations = np.sum(signature[orientation] * source_signature["normal"], axis=1)
            distance = sum((left ^ right).bit_count() for left, right in zip(signature[hash_key], source_signature["hashes"])) / len(FRACTIONS)
            orientations.append((float(correlations.mean()), -distance, orientation, float(correlations.min())))
        correlation, negative_distance, orientation, minimum = max(orientations)
        ranked.append({"source_sha256": source["sha256"], "orientation": orientation,
                       "mean_frame_correlation": round(correlation, 6), "minimum_frame_correlation": round(minimum, 6),
                       "mean_phash_hamming": round(-negative_distance, 3), "duration_error_sec": round(duration_error, 6)})
    ranked.sort(key=lambda candidate: (-candidate["mean_frame_correlation"], candidate["mean_phash_hamming"], candidate["source_sha256"]))
    for index, candidate in enumerate(ranked[:3]):
        second = ranked[index + 1]["mean_frame_correlation"] if index + 1 < len(ranked) else 0.0
        candidate["margin_to_next"] = round(candidate["mean_frame_correlation"] - second, 6)
        candidate["strong_candidate"] = (index == 0 and candidate["mean_frame_correlation"] >= 0.97
                                         and candidate["minimum_frame_correlation"] >= 0.94
                                         and candidate["mean_phash_hamming"] <= 8
                                         and candidate["margin_to_next"] >= 0.025)
    return ranked[:3]


def load_reviewed_provenance(path: Path | None, clips: list[dict], sources: list[dict]) -> dict[str, dict]:
    """Only human-attested full-video, label-preserving mappings transfer labels."""
    if path is None:
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    if value.get("schema_version") != "forte-reviewed-augmentation-provenance/1" or value.get("human_review_attested") is not True:
        raise ValueError("Augmentation provenance requires explicit human review attestation")
    if not isinstance(value.get("reviewer"), str) or not value["reviewer"].strip():
        raise ValueError("Augmentation provenance requires a reviewer")
    try:
        reviewed_at = datetime.fromisoformat(value["reviewed_at"].replace("Z", "+00:00"))
        if reviewed_at.tzinfo is None:
            raise ValueError
    except (KeyError, TypeError, AttributeError, ValueError):
        raise ValueError("Augmentation provenance requires a timezone-aware reviewed_at") from None
    known_clips, known_sources = {c["sha256"]: c for c in clips}, {s["sha256"]: s for s in sources}
    links = value.get("links")
    if not isinstance(links, list):
        raise ValueError("Reviewed provenance links must be a list")
    result = {}
    for link in links:
        if not isinstance(link, dict):
            raise ValueError("Invalid provenance link")
        clip_hash, source_hash = link.get("clip_sha256"), link.get("source_sha256")
        if (not isinstance(clip_hash, str) or not _HASH.fullmatch(clip_hash) or clip_hash not in known_clips
                or not isinstance(source_hash, str) or not _HASH.fullmatch(source_hash) or source_hash not in known_sources
                or clip_hash in result):
            raise ValueError("Provenance needs unique existing clip and source SHA-256 values")
        if link.get("label_preserving") is not True or link.get("full_video") is not True:
            raise ValueError("Binary label transfer requires reviewed full-video label preservation")
        if link.get("time_offset_sec") != 0 or link.get("time_scale") != 1:
            raise ValueError("Only unchanged full-video timelines are supported for label transfer")
        if type(link.get("horizontal_flip")) is not bool:
            raise ValueError("Provenance requires an explicit orientation")
        clip, source = known_clips[clip_hash], known_sources[source_hash]
        if clip["decode_status"] != "sampled" or source["decode_status"] != "sampled":
            raise ValueError("Reviewed sources must decode before export")
        if abs(clip["duration_sec_estimate"] - source["duration_sec_estimate"]) > max(0.25, source["duration_sec_estimate"] * 0.01):
            raise ValueError("Reviewed full-video timelines have incompatible durations")
        result[clip_hash] = {**link, "reviewer": value["reviewer"], "reviewed_at": value["reviewed_at"]}
    return result


def _component_partition(groups: list[dict], components: dict[str, set[str]], seed: int) -> dict[str, str]:
    """Keep transitive reviewed families together, including mixed-label groups."""
    if all(len(members) == 1 for members in components.values()):
        try:
            partition = deterministic_split(groups, seed)
            return {group["video_sha256"]: name for name, members in partition.items() for group in members}
        except ValueError:
            return {}
    lookup = {group["video_sha256"]: group for group in groups}
    profiles = {identity: Counter(lookup[digest]["binary_label"] for digest in members if digest in lookup)
                for identity, members in components.items()}
    total = Counter(group["binary_label"] for group in groups)
    if total[0] < 9 or total[1] < 9:
        return {}
    remaining = set(components)
    assignments = {}
    available = total.copy()
    for name in ("validation", "test"):
        counts = Counter()
        while True:
            before = sum((counts[label] - 4) ** 2 for label in (0, 1))
            choices = []
            for identity in remaining:
                profile = profiles[identity]
                if any(available[label] - profile[label] < 1 for label in (0, 1)):
                    continue
                after = sum((counts[label] + profile[label] - 4) ** 2 for label in (0, 1))
                if after < before:
                    tie = hashlib.sha256(f"{seed}:{name}:{identity}".encode()).hexdigest()
                    choices.append((after, tie, identity))
            if not choices:
                break
            identity = min(choices)[2]
            remaining.remove(identity)
            counts.update(profiles[identity])
            available.subtract(profiles[identity])
            for digest in components[identity]:
                assignments[digest] = name
        if not counts[0] or not counts[1]:
            # Insufficient separable components must not masquerade as a test.
            return {}
    for identity in remaining:
        for digest in components[identity]:
            assignments[digest] = "train"
    return assignments


def assign_supervision(sources: list[dict], clips: list[dict], labels: dict, reviewed: dict, seed: int = 17) -> dict:
    """Group identical sources, attach verified derivatives, quarantine ambiguity.

    Reviewed sources joined through any common family are unioned transitively
    before splitting. Unverified perceptual hints never transfer supervision.
    """
    rows = defaultdict(list)
    for row in labels["rows"]:
        if row["video_path"]:
            rows[row["video_path"]].append(row)
    by_hash = defaultdict(list)
    for source in sources:
        by_hash[source["sha256"]].append(source)
    groups = []
    for digest, members in sorted(by_hash.items()):
        attached = [row for member in members for row in rows[member["path"]]]
        binary = {row["binary_label"] for row in attached}
        valid = len(binary) == 1 and None not in binary and all(member["decode_status"] == "sampled" for member in members)
        groups.append({"video_sha256": digest, "binary_label": next(iter(binary)) if valid else None,
                       "video_paths": sorted(member["path"] for member in members),
                       "weak_anchors_sec": sorted({anchor for row in attached for anchor in row["weak_anchors_sec"]}),
                       "label_rows": sorted(row["row"] for row in attached), "valid": valid})
    parent = {group["video_sha256"]: group["video_sha256"] for group in groups}
    def find(digest):
        while parent[digest] != digest:
            parent[digest] = parent[parent[digest]]
            digest = parent[digest]
        return digest
    family_sources = defaultdict(set)
    for clip in clips:
        digest = clip["sha256"]
        original = digest if digest in parent else reviewed.get(digest, {}).get("source_sha256")
        if original in parent:
            family_sources[clip["family"]].add(original)
    for sources_in_family in family_sources.values():
        ordered = sorted(sources_in_family)
        for digest in ordered[1:]:
            parent[find(digest)] = find(ordered[0])
    components = defaultdict(set)
    for digest in parent:
        components[find(digest)].add(digest)
    component_ids = {digest: hashlib.sha256("|".join(sorted(members)).encode()).hexdigest()
                     for members in components.values() for digest in members}
    usable = [group for group in groups if group["valid"]]
    invalid_components = {component_ids[group["video_sha256"]] for group in groups if not group["valid"]}
    usable = [group for group in usable if component_ids[group["video_sha256"]] not in invalid_components]
    valid_hashes = {group["video_sha256"] for group in usable}
    valid_components = {identity: members for identity, members in components.items() if members <= valid_hashes}
    splits = _component_partition(usable, valid_components, seed)
    lookup = {group["video_sha256"]: group for group in groups}
    for group in groups:
        group["split"] = splits.get(group["video_sha256"], "quarantine")
        group["leakage_group_id"] = component_ids[group["video_sha256"]]
        group["supervision"] = "explicit_column_c_video_binary_only" if group["valid"] else "unknown_or_conflicting"
    families = defaultdict(list)
    for clip in clips:
        families[clip["family"]].append(clip)
        digest = clip["sha256"]
        source_hash = digest if digest in lookup else reviewed.get(digest, {}).get("source_sha256")
        clip["verified_source_sha256"] = source_hash
        source = lookup.get(source_hash)
        clip["leakage_group_id"] = source["leakage_group_id"] if source else None
        clip["binary_label"] = source["binary_label"] if source and source["valid"] else None
        clip["weak_anchors_sec"] = source["weak_anchors_sec"] if source and source["valid"] else []
        clip["split"] = splits.get(source_hash, "quarantine") if clip["decode_status"] == "sampled" else "quarantine"
        clip["provenance_status"] = "byte_identity" if digest in lookup else "human_reviewed" if digest in reviewed else "unverified"
        clip["quarantine_reason"] = None if clip["split"] != "quarantine" else "unverified_source_or_label_or_split"
    for members in families.values():
        linked = {member["verified_source_sha256"] for member in members if member["verified_source_sha256"]}
        hinted = {member["source_candidates"][0]["source_sha256"] for member in members
                  if member.get("source_candidates") and member["source_candidates"][0]["strong_candidate"]}
        family_splits = {splits[digest] for digest in linked | hinted if digest in splits}
        if len(family_splits) > 1:
            for member in members:
                member["split"] = "quarantine"
                member["quarantine_reason"] = "family_links_cross_source_splits"
    return {"source_groups": groups, "clips": clips,
            "split_policy": "transitive_verified_family_components_before_split; singleton_components_preserve_seeded_hash_holdouts; unverified_clips_quarantined",
            "camera_session_independence": "unknown; source-hash split is exploratory only"}


def export_supervised_jsonl(manifest: dict, output: Path) -> dict:
    """Export video-level binary examples only; no invented event/VLM answers."""
    output.mkdir(parents=True, exist_ok=True)
    counts = {}
    for split in ("train", "validation", "test"):
        records = []
        for source in manifest["source_groups"]:
            if source["split"] == split:
                records.append({"video_path": source["video_paths"][0], "video_sha256": source["video_sha256"],
                                "group_id": source["leakage_group_id"], "source_group_id": source["video_sha256"], "binary_label": source["binary_label"],
                                "weak_anchors_sec": source["weak_anchors_sec"], "supervision": "video_binary_only",
                                "origin": "original", "sample_weight": 1.0})
        for clip in manifest["clips"]:
            if clip["split"] == split and clip["binary_label"] is not None:
                records.append({"video_path": clip["path"], "video_sha256": clip["sha256"],
                                "group_id": clip["leakage_group_id"], "source_group_id": clip["verified_source_sha256"], "binary_label": clip["binary_label"],
                                "weak_anchors_sec": clip["weak_anchors_sec"], "supervision": "video_binary_only",
                                "origin": "reviewed_derivative" if clip["provenance_status"] == "human_reviewed" else "byte_identity",
                                "sample_weight": 1.0})
        # File copies cannot increase a video's contribution. Preserve original
        # records on byte identity. Keep clean holdouts apart from robustness
        # variants, which must be aggregated by source rather than file count.
        unique = {}
        for record in records:
            unique.setdefault(record["video_sha256"], record)
        records = list(unique.values())
        if split != "train":
            robustness = [record for record in records if record["origin"] != "original"]
            totals = Counter(record["source_group_id"] for record in robustness)
            for record in robustness:
                record["sample_weight"] = 1.0 / totals[record["source_group_id"]]
            (output / f"robustness-{split}.jsonl").write_text("".join(json.dumps(record, ensure_ascii=False) + "\n" for record in robustness), encoding="utf-8")
            records = [record for record in records if record["origin"] == "original"]
        per_group = Counter(record["source_group_id"] for record in records)
        for record in records:
            record["sample_weight"] = 1.0 / per_group[record["source_group_id"]]
        records.sort(key=lambda record: (record["group_id"], record["video_sha256"]))
        (output / f"{split}.jsonl").write_text("".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records), encoding="utf-8")
        counts[split] = len(records)
    return counts
